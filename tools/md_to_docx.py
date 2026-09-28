"""Convert the Polish reading copy (simple Markdown) to .docx.

Usage: python tools/md_to_docx.py output/IPB_Grzegorz_Piotrowski_PL.md output/IPB_Grzegorz_Piotrowski_PL.docx
"""
import re
import sys

import docx
from docx.shared import Pt


def add_runs(par, text):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if part.startswith("**") and part.endswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            par.add_run(part[1:-1]).italic = True
        elif part:
            par.add_run(part)


def main(src, dst):
    d = docx.Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for line in open(src, encoding="utf-8").read().splitlines():
        s = line.rstrip()
        if not s.strip() or s.strip() in ("---", "***"):
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            d.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
        elif re.match(r"^\s*\d+\.\s", s):
            add_runs(d.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\.\s", "", s))
        elif re.match(r"^\s*[-*]\s", s):
            add_runs(d.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s", "", s))
        elif s.startswith("|"):
            if not re.match(r"^\|[\s:|-]+\|$", s):
                add_runs(d.add_paragraph(), " | ".join(c.strip() for c in s.strip("|").split("|")))
        else:
            add_runs(d.add_paragraph(), s.strip())
    d.save(dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
